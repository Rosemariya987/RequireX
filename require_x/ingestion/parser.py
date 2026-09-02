"""
REQUIRE-X Document Ingestion & Parsing Module
Extracts textual content, document structure, and metadata from PDF, DOCX, and Plain Text SRS documents.
"""

import os
import re
from typing import Dict, Any, List, Optional
import docx
import fitz  # PyMuPDF


class DocumentParser:
    """Parses various SRS document formats into structured sections and text."""

    KNOWN_SECTION_KEYWORDS = [
        "scope", "purpose", "system overview", "overall description",
        "functional requirements", "non-functional requirements",
        "system requirements", "user requirements", "interfaces",
        "security requirements", "performance requirements",
        "reliability", "usability", "constraints", "assumptions",
        "database requirements", "architecture", "data flow",
        "use cases", "appendices", "references"
    ]

    # Phrases/terms strongly associated with genuine SRS documents (IEEE/ISO 29148 style)
    SRS_INDICATOR_KEYWORDS = [
        "software requirement specification", "srs",
        "functional requirement", "non-functional requirement",
        "system requirement", "user requirement", "stakeholder requirement",
        "acceptance criteria", "use case", "actor", "traceability",
        "requirement specification", "interface requirement",
        "performance requirement", "reliability requirement",
        "scope of the document", "overall description", "product perspective",
        "external interface requirement"
    ]

    @classmethod
    def parse_file(cls, file_path_or_buffer, filename: str) -> Dict[str, Any]:
        """
        Parses a file path or buffer based on file extension.
        Returns:
            dict with {
                'filename': str,
                'raw_text': str,
                'sections': Dict[str, str],
                'stats': Dict[str, Any]
            }
        """
        ext = os.path.splitext(filename)[1].lower()

        if ext == ".pdf":
            return cls._parse_pdf(file_path_or_buffer, filename)
        elif ext in [".docx", ".doc"]:
            return cls._parse_docx(file_path_or_buffer, filename)
        elif ext in [".txt", ".md", ".rst"]:
            return cls._parse_text(file_path_or_buffer, filename)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT, MD")

    @classmethod
    def _parse_pdf(cls, file_path_or_buffer, filename: str) -> Dict[str, Any]:
        """Parses PDF document using PyMuPDF (fitz)."""
        if isinstance(file_path_or_buffer, str):
            doc = fitz.open(file_path_or_buffer)
        else:
            # Bytes buffer (e.g. from Streamlit upload)
            doc = fitz.open(stream=file_path_or_buffer.read(), filetype="pdf")

        page_texts = []
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_texts.append(page.get_text("text"))

        raw_text = "\n\n".join(page_texts)
        stats = {
            "format": "PDF",
            "page_count": len(doc),
            "character_count": len(raw_text),
            "word_count": len(raw_text.split()),
            "lines": len(raw_text.splitlines())
        }

        sections = cls._extract_sections(raw_text)
        return {
            "filename": filename,
            "raw_text": raw_text,
            "sections": sections,
            "stats": stats
        }

    @classmethod
    def _parse_docx(cls, file_path_or_buffer, filename: str) -> Dict[str, Any]:
        """Parses DOCX document using python-docx."""
        doc = docx.Document(file_path_or_buffer)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_texts.append(row_text)

        all_text = "\n".join(paragraphs + table_texts)
        stats = {
            "format": "DOCX",
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
            "character_count": len(all_text),
            "word_count": len(all_text.split()),
            "lines": len(all_text.splitlines())
        }

        sections = cls._extract_sections(all_text)
        return {
            "filename": filename,
            "raw_text": all_text,
            "sections": sections,
            "stats": stats
        }

    @classmethod
    def _parse_text(cls, file_path_or_buffer, filename: str) -> Dict[str, Any]:
        """Parses Plain Text or Markdown document."""
        if isinstance(file_path_or_buffer, str):
            with open(file_path_or_buffer, "r", encoding="utf-8", errors="ignore") as f:
                raw_text = f.read()
        else:
            raw_text = file_path_or_buffer.read()
            if isinstance(raw_text, bytes):
                raw_text = raw_text.decode("utf-8", errors="ignore")

        stats = {
            "format": "Plain Text",
            "character_count": len(raw_text),
            "word_count": len(raw_text.split()),
            "lines": len(raw_text.splitlines())
        }

        sections = cls._extract_sections(raw_text)
        return {
            "filename": filename,
            "raw_text": raw_text,
            "sections": sections,
            "stats": stats
        }

    @classmethod
    def _extract_sections(cls, text: str) -> Dict[str, str]:
        """Identifies document sections while strictly filtering Table of Contents and page noise."""
        lines = text.splitlines()
        sections: Dict[str, List[str]] = {}
        current_section = "General / Introduction"
        sections[current_section] = []

        # Common noise patterns
        noise_pattern = re.compile(
            r"(\.\s*\.\s*\.\s*\.|\.{4,}\s*\d+|\bdepartment of computer applications\b|\bgec thrissur\b|^\s*\d+\s*$|^chapter\s+\d+|^\s*page\s+\d+\s*$)",
            re.IGNORECASE
        )

        toc_mode = False

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                continue

            # Check if entering Table of Contents
            if re.match(r"^(?:table of\s+)?contents|list of figures|list of tables|list of abbreviations", trimmed, re.IGNORECASE):
                toc_mode = True
                continue

            # If in TOC mode and see dots or page references, ignore
            if toc_mode:
                if re.search(r"\.\s*\.\s*\.\s*\.|\.{3,}\s*\d+", trimmed) or len(trimmed) < 40:
                    continue
                # If we encounter a real chapter or major section header without dots, exit TOC mode
                if re.match(r"^(?:chapter\s+\d+|1\s+introduction|4\s+system analysis)", trimmed, re.IGNORECASE):
                    toc_mode = False

            # Ignore noise lines
            if noise_pattern.search(trimmed):
                continue

            # Check if this line looks like a major section header
            is_header = False
            lower_trimmed = trimmed.lower()

            # Check for header format e.g. "1. Functional Requirements", "4.6 Functional Requirements" or "# Functional Requirements"
            if not re.search(r"\.{3,}", trimmed):
                if trimmed.startswith(("#", "==")) or re.match(r"^(?:#+\s*|\d+(?:\.\d+)*\.?\s+)[A-Z]", trimmed) or any(k in lower_trimmed for k in cls.KNOWN_SECTION_KEYWORDS):
                    clean_header = re.sub(r"^[#\=\*\-\d\.\s]+", "", trimmed).rstrip(":").strip()
                    if any(k in clean_header.lower() for k in cls.KNOWN_SECTION_KEYWORDS) or len(clean_header.split()) <= 6:
                        if 3 <= len(clean_header) <= 60:
                            current_section = clean_header
                            if current_section not in sections:
                                sections[current_section] = []
                            is_header = True

            if not is_header:
                sections[current_section].append(trimmed)

        # Convert lists back to string content
        return {sec: "\n".join(content) for sec, content in sections.items() if content}

    @classmethod
    def assess_srs_likelihood(cls, raw_text: str, sections: Dict[str, str]) -> Dict[str, Any]:
        """
        Heuristically assesses whether a parsed document actually resembles an SRS,
        so the pipeline can warn the user instead of silently analyzing an unrelated file.

        Returns:
            dict with {
                'is_srs': bool,
                'confidence': float (0.0 - 1.0),
                'matched_keywords': List[str],
                'matched_sections': List[str],
                'shall_statement_count': int,
                'reason': str
            }
        """
        text = (raw_text or "").strip()

        if len(text) < 100:
            return {
                "is_srs": False,
                "confidence": 0.0,
                "matched_keywords": [],
                "matched_sections": [],
                "shall_statement_count": 0,
                "reason": "The document is too short or empty to contain a meaningful SRS."
            }

        text_lower = text.lower()
        section_titles_lower = " | ".join(sections.keys()).lower()

        matched_keywords = [kw for kw in cls.SRS_INDICATOR_KEYWORDS if kw in text_lower]
        matched_sections = [kw for kw in cls.KNOWN_SECTION_KEYWORDS if kw in section_titles_lower]
        shall_count = len(re.findall(r"\bshall\b", text_lower))

        # Weighted heuristic score out of 100
        score = 0
        score += min(len(matched_keywords), 6) * 8      # up to 48
        score += min(len(matched_sections), 4) * 8       # up to 32
        score += min(shall_count, 10) * 2                # up to 20

        confidence = round(min(score / 100, 1.0), 2)
        is_srs = score >= 35

        if is_srs:
            reason = "Document contains SRS-style terminology, requirement sections, and/or 'shall' statements."
        else:
            reason = (
                "Document does not contain typical SRS terminology (e.g. functional/non-functional "
                "requirements, scope, use cases, 'shall' statements). It may be a resume, report, "
                "code file, or unrelated document."
            )

        return {
            "is_srs": is_srs,
            "confidence": confidence,
            "matched_keywords": matched_keywords,
            "matched_sections": matched_sections,
            "shall_statement_count": shall_count,
            "reason": reason
        }
