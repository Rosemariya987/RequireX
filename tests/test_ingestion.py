"""
Tests for Ingestion & Preprocessing Modules (PDF, DOCX, TXT)
"""
import io
import os
import unittest
import docx
import fitz
from require_x.ingestion.parser import DocumentParser
from require_x.ingestion.preprocessor import TextPreprocessor


class TestIngestion(unittest.TestCase):
    def test_text_parser(self):
        sample_text = """
# 1. System Overview
SmartCare is a healthcare application.

# 2. Functional Requirements
* FR-001: The system shall register patients.
* FR-002: The system shall authenticate doctors.
"""
        buffer = io.StringIO(sample_text)
        result = DocumentParser.parse_file(buffer, "sample.txt")
        self.assertEqual(result["stats"]["format"], "Plain Text")
        self.assertTrue(len(result["sections"]) > 0)

        candidates = TextPreprocessor.preprocess_text(result["raw_text"], result["sections"])
        self.assertGreaterEqual(len(candidates), 2)

    def test_docx_parser(self):
        # Create an in-memory docx
        doc = docx.Document()
        doc.add_heading("1. Functional Requirements", level=1)
        doc.add_paragraph("FR-001: The system shall process online payments.")
        doc.add_paragraph("FR-002: The system shall send transaction receipts.")
        
        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)

        result = DocumentParser.parse_file(bio, "test_srs.docx")
        self.assertEqual(result["stats"]["format"], "DOCX")
        self.assertIn("Functional Requirements", result["sections"])
        self.assertIn("FR-001", result["raw_text"])

    def test_pdf_parser(self):
        # Create an in-memory PDF using fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "1. Non-Functional Requirements\nNFR-001: The system shall encrypt all sensitive patient data using AES-256.")
        
        pdf_bytes = doc.tobytes()
        doc.close()
        bio = io.BytesIO(pdf_bytes)

        result = DocumentParser.parse_file(bio, "test_srs.pdf")
        self.assertEqual(result["stats"]["format"], "PDF")
        self.assertEqual(result["stats"]["page_count"], 1)
        self.assertIn("NFR-001", result["raw_text"])


if __name__ == "__main__":
    unittest.main()
